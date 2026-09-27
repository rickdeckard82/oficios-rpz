import re
import unicodedata
import zipfile
from datetime import datetime
from html import unescape
from html.parser import HTMLParser
from ipaddress import ip_address, ip_interface
from pathlib import Path
from xml.etree import ElementTree
from urllib.parse import urlparse

import openpyxl
import xlrd
from pypdf import PdfReader

DOMAIN_RE = re.compile(
    r"(?:https?://)?(?:[a-z0-9](?:[a-z0-9-]{0,61}[a-z0-9])?\.)+"
    r"(?:[a-z]{2,63}|xn--[a-z0-9-]{1,59}[a-z0-9])(?:/[^\s,;]*)?",
    re.IGNORECASE,
)
EMAIL_RE = re.compile(
    r"\b[a-z0-9._%+-]+@(?:[a-z0-9](?:[a-z0-9-]{0,61}[a-z0-9])?\.)+[a-z]{2,63}\b",
    re.IGNORECASE,
)
IP_CANDIDATE_RE = re.compile(r"(?<![a-zA-Z0-9_.:/-])(?:[0-9]{1,3}(?:\.[0-9]{1,3}){3}(?:/\d{1,2})?|[0-9a-fA-F:]{2,}(?:/\d{1,3})?)(?![a-zA-Z0-9_.:/-])")
HTML_CHARSET_RE = re.compile(rb"charset=([a-zA-Z0-9_-]+)")
SIGNATURE_DATE_RE = re.compile(r"documento assinado eletronicamente.*?\bem\s+(\d{2}/\d{2}/\d{4})", re.IGNORECASE | re.DOTALL)
SEI_NUMBER_RE = re.compile(r"SEI\s*n[ºo]?\s*[:\s]*([0-9]{5,})", re.IGNORECASE)
OFFICE_NUMBER_RE = re.compile(
    "\\bOf(?:i|\\u00ed)cio\\s*n(?:[.\\u00ba]|o)?\\s*[:\\-]?\\s*"
    "([0-9]{1,6}\\s*/\\s*[0-9]{4}(?:\\s*/\\s*[A-Z0-9-]+)+)",
    re.IGNORECASE,
)

MULTI_PART_SUFFIXES = {
    "com.br",
    "net.br",
    "org.br",
    "gov.br",
    "edu.br",
    "mil.br",
    "jus.br",
    "leg.br",
    "blog.br",
    "tv.br",
}
MULTI_PART_SUFFIX_PREFIXES = {suffix.split(".", 1)[0] for suffix in MULTI_PART_SUFFIXES}
ODS_NAMESPACES = {
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
}


def extract_domains(path):
    suffix = Path(path).suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return extract_xlsx_domains(path)
    if suffix == ".xls":
        return extract_xls_domains(path)
    if suffix == ".ods":
        return extract_ods_domains(path)

    domains = set()
    for text in iter_extracted_text(path):
        domains.update(extract_domains_from_text(text))
    return sorted(domain for domain in domains if domain)


def extract_ip_addresses(path):
    addresses = set()
    for text in iter_extracted_text(path):
        for match in IP_CANDIDATE_RE.finditer(text):
            normalized = normalize_ip_address(match.group(0))
            if normalized:
                addresses.add(normalized)
    return sorted(addresses, key=ip_sort_key)


def extract_office_metadata(path):
    text = extract_document_text(path)
    return {
        "office_number": extract_office_number(text),
        "sei_numbers": extract_annex_sei_numbers(text),
        "expedition_date": extract_signature_date(text),
    }


def extract_document_text(path):
    suffix = Path(path).suffix.lower()
    if suffix in {".html", ".htm"}:
        return extract_html_text(path)
    if suffix == ".pdf":
        return "\n".join(extract_pdf_text(path))
    return ""


def extract_html_text(path):
    raw = Path(path).read_bytes()
    charset_match = HTML_CHARSET_RE.search(raw[:4096])
    encodings = []
    if charset_match:
        encodings.append(charset_match.group(1).decode("ascii", errors="ignore"))
    encodings.extend(["utf-8", "iso-8859-1", "latin-1"])

    for encoding in encodings:
        try:
            html = raw.decode(encoding)
            break
        except (LookupError, UnicodeDecodeError):
            continue
    else:
        html = raw.decode("utf-8", errors="replace")

    parser = TextHTMLParser()
    parser.feed(html)
    return parser.text()


def extract_annex_sei_numbers(text):
    annex_text = text_after_label(text, "Anexos:")
    if not annex_text:
        annex_text = text_after_label(text, "Anexo:")
    if annex_text:
        annex_text = annex_text.split("Atenciosamente", 1)[0]
    else:
        annex_text = text

    numbers = []
    for match in SEI_NUMBER_RE.finditer(annex_text):
        number = match.group(1)
        if number not in numbers:
            numbers.append(number)
    return numbers


def extract_office_number(text):
    match = OFFICE_NUMBER_RE.search(" ".join(text.split()))
    if not match:
        return ""
    return re.sub(r"\s*/\s*", "/", match.group(1).strip().strip(".,;:"))


def extract_signature_date(text):
    normalized = " ".join(text.split())
    match = SIGNATURE_DATE_RE.search(normalized)
    if not match:
        return None
    try:
        return datetime.strptime(match.group(1), "%d/%m/%Y").date()
    except ValueError:
        return None


def text_after_label(text, label):
    normalized_label = normalize_header(label)
    lines = text.splitlines()
    for index, line in enumerate(lines):
        if normalize_header(line).startswith(normalized_label):
            return "\n".join(lines[index:])
    return ""


class TextHTMLParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []

    def handle_data(self, data):
        text = unescape(data).strip()
        if text:
            self.parts.append(text)

    def handle_starttag(self, tag, attrs):
        if tag in {"br", "p", "tr", "td", "div", "table"}:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"p", "tr", "div", "table"}:
            self.parts.append("\n")

    def text(self):
        lines = []
        for line in "".join(self.parts).splitlines():
            clean = " ".join(line.split())
            if clean:
                lines.append(clean)
        return "\n".join(lines)


def normalize_ip_address(value):
    cleaned = value.strip().strip(".,;:()[]{}<>\"'")
    if not cleaned:
        return ""

    try:
        if "/" in cleaned:
            interface = ip_interface(cleaned)
            if interface.version == 4 and interface.network.prefixlen == 32:
                return str(interface.ip)
            if interface.version == 6 and interface.network.prefixlen == 128:
                return f"{interface.ip}/128"
            return f"{interface.ip}/{interface.network.prefixlen}"

        address = ip_address(cleaned)
        if address.version == 4:
            return str(address)
        return f"{address}/128"
    except ValueError:
        return ""


def ip_sort_key(value):
    address = value.split("/", 1)[0]
    parsed = ip_address(address)
    return (parsed.version, int(parsed))


def remove_email_addresses(value):
    return EMAIL_RE.sub(" ", value)


def iter_extracted_text(path):
    suffix = Path(path).suffix.lower()
    if suffix == ".pdf":
        yield from extract_pdf_text(path)
    elif suffix in {".xlsx", ".xlsm"}:
        yield from extract_xlsx_text(path)
    elif suffix == ".xls":
        yield from extract_xls_text(path)
    elif suffix == ".ods":
        yield from extract_ods_text(path)
    else:
        raise ValueError("Formato nao suportado. Envie PDF, XLS, XLSX ou ODS.")


def extract_pdf_text(path):
    reader = PdfReader(path)
    for page in reader.pages:
        text = page.extract_text() or ""
        yield from text.splitlines()


def extract_xlsx_text(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for cell in row:
                    if cell is not None:
                        yield str(cell)
    finally:
        workbook.close()


def extract_xlsx_domains(path):
    domains = set()
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            domains.update(extract_domains_from_rows(rows))
    finally:
        workbook.close()
    return sorted(domain for domain in domains if domain)


def extract_xls_text(path):
    workbook = xlrd.open_workbook(path)
    for sheet in workbook.sheets():
        for row_idx in range(sheet.nrows):
            for value in sheet.row_values(row_idx):
                if value:
                    yield str(value)


def extract_xls_domains(path):
    domains = set()
    workbook = xlrd.open_workbook(path)
    for sheet in workbook.sheets():
        rows = [sheet.row_values(row_idx) for row_idx in range(sheet.nrows)]
        domains.update(extract_domains_from_rows(rows))
    return sorted(domain for domain in domains if domain)


def extract_ods_text(path):
    for row in extract_ods_rows(path):
        for value in row:
            if value is not None:
                yield str(value)


def extract_ods_domains(path):
    domains = set()
    domains.update(extract_domains_from_rows(extract_ods_rows(path)))
    return sorted(domain for domain in domains if domain)


def extract_ods_rows(path):
    with zipfile.ZipFile(path) as ods_file:
        content = ods_file.read("content.xml")

    root = ElementTree.fromstring(content)
    rows = []
    for table in root.findall(".//table:table", ODS_NAMESPACES):
        for row in table.findall("table:table-row", ODS_NAMESPACES):
            repeat_rows = ods_repeat_count(row, "number-rows-repeated")
            values = []
            for cell in row.findall("table:table-cell", ODS_NAMESPACES):
                repeat_columns = ods_repeat_count(cell, "number-columns-repeated")
                text = ods_cell_text(cell)
                values.extend([text] * repeat_columns)
            if any(value for value in values):
                rows.extend([values] * repeat_rows)
    return rows


def ods_repeat_count(element, attribute_name):
    value = element.get(f"{{{ODS_NAMESPACES['table']}}}{attribute_name}")
    if not value:
        return 1
    try:
        return min(int(value), 1000)
    except ValueError:
        return 1


def ods_cell_text(cell):
    parts = []
    for paragraph in cell.findall(".//text:p", ODS_NAMESPACES):
        text = "".join(paragraph.itertext()).strip()
        if text:
            parts.append(text)
    return "\n".join(parts)


def extract_domains_from_rows(rows):
    domains = set()
    header_idx, new_domain_columns = find_new_domain_columns(rows)
    if new_domain_columns:
        data_rows = rows[header_idx + 1 :]
        for row in data_rows:
            for column_idx in new_domain_columns:
                if column_idx < len(row) and row[column_idx] is not None:
                    domains.update(extract_domains_from_text(row[column_idx]))
        return domains

    for row in rows:
        for value in row:
            if value is not None:
                domains.update(extract_domains_from_text(value))
    return domains


def find_new_domain_columns(rows):
    for row_idx, row in enumerate(rows[:20]):
        columns = [
            column_idx
            for column_idx, value in enumerate(row)
            if is_new_domain_header(value)
        ]
        if columns:
            return row_idx, columns
    return None, []


def is_new_domain_header(value):
    normalized = normalize_header(value)
    return "dominio" in normalized and ("novo" in normalized or "novos" in normalized)


def normalize_header(value):
    text = unicodedata.normalize("NFKD", str(value or "").strip().lower())
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def extract_domains_from_text(value):
    text = remove_email_addresses(str(value))
    domains = set()
    for match in DOMAIN_RE.finditer(text):
        domain = normalize_domain(match.group(0))
        if domain:
            domains.add(domain)
    return domains


# Cada nome vira "*.<dominio>.rpz.zone." na zona RPZ (sem ponto final, herda o $ORIGIN);
# acima disso o nome final estoura os 255 bytes do formato wire do DNS e derruba o load da zona inteira no BIND.
MAX_RPZ_DOMAIN_LENGTH = 242


def normalize_domain(value):
    cleaned = value.strip().lower().strip(".,;:()[]{}<>\"'")
    if "://" not in cleaned:
        cleaned = f"//{cleaned}"

    parsed = urlparse(cleaned)
    domain = (parsed.hostname or parsed.path.split("/", 1)[0]).strip(".")
    if domain.startswith("www."):
        domain = domain[4:]

    labels = [label for label in domain.split(".") if label]
    if len(labels) < 2:
        return ""

    suffix = ".".join(labels[-2:])
    if suffix in MULTI_PART_SUFFIXES and len(labels) == 2:
        return ""

    if not is_valid_domain_labels(labels):
        return ""

    result = ".".join(labels)
    if len(result) > MAX_RPZ_DOMAIN_LENGTH:
        return ""

    return result


def is_valid_domain_labels(labels):
    if any(not re.fullmatch(r"[a-z0-9](?:[a-z0-9-]{0,61}[a-z0-9])?", label) for label in labels):
        return False

    return is_valid_tld(labels[-1])


def is_valid_tld(tld):
    if re.fullmatch(r"[a-z]{2,63}", tld):
        return True
    return bool(re.fullmatch(r"xn--[a-z0-9-]{1,59}[a-z0-9]", tld))
